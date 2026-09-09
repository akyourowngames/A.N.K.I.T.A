"""Evidence integrity and persistence tests. Models are isolated test doubles."""
import json
import sqlite3
import time
from io import BytesIO

import pytest
from knowledge import extraction, parsers, service, storage


@pytest.fixture(autouse=True)
def isolated(tmp_path, monkeypatch):
    monkeypatch.setenv("ZUMBA_KNOWLEDGE_HOME", str(tmp_path / "graph"))
    monkeypatch.setenv("ZUMBA_MEMORY_HOME", str(tmp_path / "memory"))
    monkeypatch.setattr(service.embedder, "embed_texts", lambda texts: [[1., 0., 0.] for _ in texts])
    monkeypatch.setattr(service.embedder, "embed_text", lambda text: [1., 0., 0.])
    monkeypatch.setattr(service.llm, "_memory_model", lambda: "test-model")
    monkeypatch.setattr(service.llm, "chat_json", lambda *a, **k: {"id": None, "reason": "Distinct identity"})


def proposed():
    return {"entities": [
        {"name": "Mira", "type": "Person", "evidence": "Mira founded Atlas.", "confidence": .94},
        {"name": "Atlas", "type": "Company", "evidence": "Mira founded Atlas.", "confidence": .93}],
        "relations": [{"source": "Mira", "target": "Atlas", "type": "FOUNDED", "evidence": "Mira founded Atlas.", "confidence": .92}]}


def test_quotes_and_semantic_verification_fail_closed(monkeypatch):
    result = proposed()
    result["relations"].append({**result["relations"][0], "type": "OWNS", "evidence": "Mira owns Atlas."})
    responses = iter([result, {"entities": [0, 1], "relations": [0]}])
    monkeypatch.setattr(extraction.llm, "chat_json", lambda *a, **k: next(responses))
    entities, edges, rejected = extraction.extract_graph("Mira founded Atlas.", "txt")
    assert len(entities) == 2 and len(edges) == 1 and rejected == 1
    assert edges[0]["type"] == "FOUNDED"


def test_verifier_rejects_semantically_wrong_but_real_quote(monkeypatch):
    result = proposed()
    result["relations"][0]["type"] = "ACQUIRED"
    responses = iter([result, {"entities": [0, 1], "relations": []}])
    monkeypatch.setattr(extraction.llm, "chat_json", lambda *a, **k: next(responses))
    _, edges, rejected = extraction.extract_graph("Mira founded Atlas.", "txt")
    assert edges == [] and rejected == 1


def test_invalid_verification_never_saves_graph(monkeypatch):
    responses = iter([proposed(), None])
    monkeypatch.setattr(extraction.llm, "chat_json", lambda *a, **k: next(responses))
    doc_id = service.enqueue("report.txt", "txt", b"Mira founded Atlas.")
    service.ingest(doc_id)
    graph = service.snapshot()
    assert not graph["nodes"] and not graph["edges"]
    assert graph["documents"][0]["status"] == "failed"
    assert graph["counts"]["chunks"] == 1


def test_ingest_provenance_dedup_and_restart(monkeypatch):
    result = proposed()
    monkeypatch.setattr(extraction, "extract_graph", lambda *a: (result["entities"], result["relations"], 0))
    doc_id = service.enqueue("report.md", "md", b"# Founding\nMira founded Atlas.")
    assert service.enqueue("renamed.md", "md", b"# Founding\nMira founded Atlas.") == doc_id
    service.ingest(doc_id)
    graph = service.snapshot()
    assert len(graph["nodes"]) == 2 and len(graph["edges"]) == 1
    ev = graph["edges"][0]["evidence"][0]
    assert ev["document"] == "report.md" and ev["section"] == "Founding"
    assert ev["quote"] == "Mira founded Atlas." and ev["page"] is None
    assert ev["document_id"] == doc_id and ev["method"].startswith("test-model")
    assert graph["counts"] == {"chunks": 1, "embedded": 1}
    # Complete chunks survive reopening the database and are not re-extracted.
    monkeypatch.setattr(extraction, "extract_graph", lambda *a: pytest.fail("Already complete chunk re-extracted"))
    service.ingest(doc_id)
    assert len(service.snapshot()["edges"]) == 1


def test_resume_partial_document_without_duplicating_edges(monkeypatch):
    monkeypatch.setattr(parsers, "chunk", lambda sections: iter([parsers.Section("Mira founded Atlas.", "One"), parsers.Section("Second section", "Two")]))
    result = proposed()
    def extract(text, kind):
        if text == "Second section":
            raise RuntimeError("Temporary model failure")
        return result["entities"], result["relations"], 0
    monkeypatch.setattr(extraction, "extract_graph", extract)
    doc_id = service.enqueue("report.txt", "txt", b"Mira founded Atlas. Second section")
    service.ingest(doc_id)
    assert service.snapshot()["documents"][0]["status"] == "failed"
    seen = []
    monkeypatch.setattr(extraction, "extract_graph", lambda text, kind: (seen.append(text) or [], [], 0))
    service.ingest(doc_id)
    assert seen == ["Second section"]
    graph = service.snapshot()
    assert graph["documents"][0]["status"] == "complete"
    assert len(graph["edges"]) == 1 and len(graph["edges"][0]["evidence"]) == 1


def test_entity_resolution_requires_reasoning_even_identical_vectors(monkeypatch):
    result = proposed()
    monkeypatch.setattr(extraction, "extract_graph", lambda *a: (result["entities"], result["relations"], 0))
    first = service.enqueue("one.txt", "txt", b"Mira founded Atlas.")
    service.ingest(first)
    graph = service.snapshot()
    lookup = {n["name"]: n["id"] for n in graph["nodes"]}
    def merge(prompt, **kwargs):
        data = json.loads(prompt.split("\n", 1)[1])
        return {"id": lookup[data["new"]["name"]], "reason": "Same founding event in context"}
    monkeypatch.setattr(service.llm, "chat_json", merge)
    second = service.enqueue("two.txt", "txt", b"Mira founded Atlas.\n")
    service.ingest(second)
    graph = service.snapshot()
    assert len(graph["nodes"]) == 2 and len(graph["edges"]) == 1
    assert graph["edges"][0]["sourceCount"] == 2
    assert len([a for a in graph["audit"] if a["decision"] == "merged"]) == 2


def test_memory_sync_reads_only_user_evidence_and_versions(tmp_path):
    root = tmp_path / "memory"
    root.mkdir()
    (root / "user.md").write_text("# Profile\nMira founded Atlas.", encoding="utf-8")
    with sqlite3.connect(root / "memory.db") as con:
        con.execute("CREATE TABLE episodes(id INTEGER PRIMARY KEY,user_text TEXT,assistant_text TEXT,session_id TEXT)")
        con.execute("INSERT INTO episodes VALUES(1,'Mira founded Atlas.','Mira owns the moon.','s')")
    service.sync_sources()
    service.sync_sources()
    with storage.connect() as con:
        assert con.execute("SELECT COUNT(*) FROM documents").fetchone()[0] == 2
        assert all(b"moon" not in row[0] for row in con.execute("SELECT content FROM documents"))
    (root / "user.md").write_text("# Profile\nUpdated profile", encoding="utf-8")
    service.sync_sources()
    assert len(service.snapshot()["documents"]) == 3


@pytest.mark.parametrize("kind", ["txt", "md", "markdown", "json", "csv"])
def test_text_parsers_preserve_original_content(kind):
    text = '{"company":"Atlas","founder":"Mira"}' if kind == "json" else 'name,statement\nMira,"founded Atlas, Inc."\n' if kind == "csv" else "# Report\nMira founded Atlas.\n"
    sections = parsers.extract(text.encode(), kind)
    assert "".join(s.text for s in sections) == text


def test_pdf_has_actual_page_numbers():
    import fitz
    with fitz.open() as doc:
        doc.new_page().insert_text((72, 72), "First page.")
        doc.new_page().insert_text((72, 72), "Mira founded Atlas.")
        content = doc.tobytes()
    sections = parsers.extract(content, "pdf")
    assert sections[1].page == 2 and "Mira founded Atlas." in sections[1].text


def test_docx_has_sections_not_invented_pages():
    from docx import Document
    doc = Document()
    doc.add_heading("Founding", 1)
    doc.add_paragraph("Mira founded Atlas.")
    stream = BytesIO(); doc.save(stream)
    sections = parsers.extract(stream.getvalue(), "docx")
    assert sections[-1].section == "Founding" and sections[-1].page is None


def test_corrupt_and_empty_documents_fail():
    for content, kind in [(b"{broken", "json"), (b"", "txt"), (b"not a zip", "docx")]:
        with pytest.raises(Exception):
            parsers.extract(content, kind)


def test_chunking_covers_text_and_preserves_page():
    text = "Mira founded Atlas.\n" * 600
    chunks = list(parsers.chunk([parsers.Section(text, "Page 4", 4)]))
    assert all(c.page == 4 and c.text in text and len(c.text) <= 4800 for c in chunks)
    assert chunks[-1].text.endswith("Mira founded Atlas.\n")


def test_api_upload_retry_and_download(monkeypatch):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from server.knowledge_api import router
    app = FastAPI(); app.include_router(router)
    client = TestClient(app)
    response = client.post("/api/knowledge/documents", files={"file": ("../report.txt", b"Mira founded Atlas.")})
    assert response.status_code == 202
    doc_id = response.json()["id"]
    assert client.get(f"/api/knowledge/documents/{doc_id}/download").content == b"Mira founded Atlas."
    assert client.post(f"/api/knowledge/documents/{doc_id}/retry").status_code == 409
    assert client.post("/api/knowledge/documents", files={"file": ("bad.exe", b"x")}).status_code == 415
    assert client.get("/api/knowledge/documents/unknown").status_code == 404
    assert client.get("/api/knowledge").json()["documents"][0]["name"] == "report.txt"


def test_unverifiable_answer_is_not_returned(monkeypatch):
    from server.knowledge_api import query, Query
    doc_id = service.enqueue("report.txt", "txt", b"Mira founded Atlas.")
    monkeypatch.setattr(extraction, "extract_graph", lambda *a: ([], [], 0))
    service.ingest(doc_id)
    responses = iter([{"search": "Mira"}, {"answer": "Mira owns the moon.", "citations": [{"index": 1, "quote": "Mira owns the moon."}]}])
    monkeypatch.setattr(service.llm, "chat_json", lambda *a, **k: next(responses))
    result = query(Query(question="Who is Mira?"))
    assert not result["citations"] and "owns the moon" not in result["answer"]


def test_durable_inbox_survives_reopen(tmp_path, monkeypatch):
    from memory import inbox
    monkeypatch.setattr(inbox.db, "memory_home", lambda: tmp_path)
    ident = inbox.save("Mira founded Atlas.", "Saved", "s", "chat")
    assert inbox.save("Mira founded Atlas.", "Saved", "s", "chat") == ident
    assert inbox.pending() == [(ident, "Mira founded Atlas.", "Saved", "s", "chat")]
    inbox.finish(ident)
    assert inbox.pending() == []


def test_flush_respects_timeout():
    from memory.service import Memory
    memory = Memory(con=object())
    memory._idle.clear()
    start = time.monotonic()
    assert memory.flush(.02) is False
    assert time.monotonic() - start < .5


def test_profile_selection_is_llm_driven(monkeypatch):
    from identity import userprofile
    from memory import llm as memory_llm
    captured = []
    # identity.userprofile calls memory.llm directly, so patch that module —
    # patching knowledge.reasoning would leave a live network call in the test.
    monkeypatch.setattr(memory_llm, "chat_json", lambda *a, **k: {"facts": [{"index": 1, "key": "user_project"}]})
    monkeypatch.setattr(userprofile, "upsert_fact", lambda con, key, value, conf, ep: captured.append((key, value)))
    userprofile.extract_user_facts_from_relations(None, [{"fact": "Mira likes tea"}, {"fact": "The user founded Atlas", "type": "custom_relation"}])
    assert captured == [("user_project", "The user founded Atlas")]


def test_bulk_retry_only_requeues_failed_sources():
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from server.knowledge_api import router
    failed = service.enqueue("failed.txt", "txt", b"one")
    complete = service.enqueue("complete.txt", "txt", b"two")
    processing = service.enqueue("processing.txt", "txt", b"three")
    service.stage(failed, "Needs attention", status="failed", error="provider failure")
    service.stage(complete, "Complete", status="complete")
    service.stage(processing, "Extracting", status="processing")
    app = FastAPI(); app.include_router(router)
    response = TestClient(app).post("/api/knowledge/documents/retry-failed")
    assert response.status_code == 200
    assert response.json()["queued"] == 1
    states = {d["id"]: d["status"] for d in service.snapshot()["documents"]}
    assert states == {failed: "queued", complete: "complete", processing: "processing"}
